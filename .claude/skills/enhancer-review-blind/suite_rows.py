"""Read-only dump of a suite SNAPSHOT for enhancer-review-blind.

Usage:
    python suite_rows.py "<run_dir>\\snapshot_<name>.xlsx|.md" [output.txt]

.xlsx: prints every non-empty row of every sheet with its row number, section and
row kind (SECTION / HEADER / TC / OTHER).
.md:   prints every line with its line number, marking TC headings and sections.

Refuses any file that is not a snapshot inside the enhancer runs folder, so the
blind pass cannot read the live (after-write) suite through this script.
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

RUNS_DIR = Path(r"C:\Users\pkatakam\.claude\skills\jira-test-enhancer\runs")
MAIN_SHEET = "CRA Regression Suite"
HEADER_NAMES = ["TC Name", "Scenario", "Stay in Regression", "Test Steps", "Expected Result"]


def check_is_snapshot(path: Path) -> None:
    resolved = path.resolve()
    if (
        not resolved.name.startswith("snapshot_")
        or resolved.suffix.lower() not in (".xlsx", ".md")
        or RUNS_DIR.resolve() not in resolved.parents
    ):
        sys.exit(
            f"REFUSED: {resolved} is not a snapshot inside {RUNS_DIR}. "
            "The blind pass may only read snapshot_*.xlsx / snapshot_*.md files in a run folder."
        )


def dump_md(path: Path) -> str:
    lines = [f"Snapshot: {path}", ""]
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        stripped = line.lstrip("#").strip()
        if line.startswith("#") and stripped.startswith("TC-"):
            marker = "TC     "
        elif line.startswith("#"):
            marker = "SECTION"
        else:
            marker = "       "
        lines.append(f"[line {line_number:>4}] {marker} {line}")
    return "\n".join(lines)


def dump(path: Path) -> str:
    if path.suffix.lower() == ".md":
        return dump_md(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines = [f"Snapshot: {path}", f"Sheets: {' | '.join(workbook.sheetnames)}", ""]

    for name in workbook.sheetnames:
        sheet = workbook[name]
        lines.append(f"=== Sheet: {name} ===")
        section = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if v is None else str(v).strip() for v in row[: len(HEADER_NAMES)]]
            non_empty = [v for v in values if v]
            if not non_empty:
                continue
            if name == MAIN_SHEET and values == HEADER_NAMES:
                kind = "HEADER"
            elif name == MAIN_SHEET and "_Regr_" in values[0]:
                kind = "TC"
            elif len(non_empty) == 1 and values[0]:
                kind, section = "SECTION", values[0]
            else:
                kind = "OTHER"
            lines.append(f"[row {row_number}] {kind} | section: {section}")
            for header, value in zip(HEADER_NAMES, values):
                if value:
                    lines.append(f"    {header if name == MAIN_SHEET else 'col'}: {value}")
        lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    snapshot = Path(sys.argv[1])
    if not snapshot.exists():
        sys.exit(f"ERROR: file not found: {snapshot}")
    check_is_snapshot(snapshot)
    report = dump(snapshot)
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(report, encoding="utf-8")
        print(f"written: {sys.argv[2]}")
    else:
        print(report)
