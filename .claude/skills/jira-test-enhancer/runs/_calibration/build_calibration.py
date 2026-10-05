"""Builds the enhancer-review calibration cases from a frozen copy of the BST suite.

Usage:
    python build_calibration.py

Creates (overwriting previous outputs, never touching the live suite):
    base_suite.xlsx                      frozen copy of the live suite (made once, never refreshed)
    seeded\\snapshot_BST CRA Regression Suite.xlsx   "before" = base minus the DAP-1654 section
    seeded\\change-summary.md
    control\\snapshot_BST CRA Regression Suite.xlsx  same "before"
    control\\change-summary.md
    live\\seeded.xlsx                     "after" with one planted mistake per check R1-R10
    live\\control.xlsx                    "after" = base (the real DAP-1654 enhancement)
    answer-key.md                        what was planted where, and the pass rule

Source for both cases: the DAP-1654 E2E file only (tickets: none), so the answer
key can't drift when the Jira ticket is edited.
"""
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
LIVE_SUITE = Path(r"C:\2026\Fall 2026\Test Cases\BST CRA Regression Suite.xlsx")
E2E_FILE = Path(r"C:\Users\pkatakam\.claude\skills\jira-e2e-test-writer\Fall26\TestCases_DAP-1654_Fall26.md")
BASE = HERE / "base_suite.xlsx"
SNAPSHOT_NAME = "snapshot_BST CRA Regression Suite.xlsx"
SHEET = "CRA Regression Suite"

SECTION_FIRST_ROW = 147        # blank separator before the DAP-1654 section title
HEADER_ROW = 149
FIRST_TC_ROW, LAST_TC_ROW = 150, 163
FIRST_NUMBER = 124
PREFIX = "CRA_TempcredsMR_Regr_"

# Real section: TC number -> the E2E TC(s) it came from.
REAL_SOURCES = {
    124: "E2E TC-DAP-1654-001", 125: "E2E TC-DAP-1654-004", 126: "E2E TC-DAP-1654-005",
    127: "E2E TC-DAP-1654-006", 128: "E2E TC-DAP-1654-007", 129: "E2E TC-DAP-1654-008",
    130: "E2E TC-DAP-1654-009", 131: "E2E TC-DAP-1654-010, -011", 132: "E2E TC-DAP-1654-012, -013",
    133: "E2E TC-DAP-1654-017, -018", 134: "E2E TC-DAP-1654-021", 135: "E2E TC-DAP-1654-027",
    136: "E2E TC-DAP-1654-030", 137: "E2E TC-DAP-1654-020",
}
EXCLUSIONS = {
    "E2E TC-DAP-1654-002": "RAOTT cache internals; not a regression-level handoff",
    "E2E TC-DAP-1654-003": "RAOTT cache internals; not a regression-level handoff",
    "E2E TC-DAP-1654-014": "lockout timing; confirm 5th-attempt threshold with dev team first",
    "E2E TC-DAP-1654-015": "lockout progression timing (deep edge)",
    "E2E TC-DAP-1654-016": "auto-unlock timing (deep edge)",
    "E2E TC-DAP-1654-019": "405 method check is cosmetic (P3)",
    "E2E TC-DAP-1654-022": "WAF injection is a one-off security scan",
    "E2E TC-DAP-1654-023": "input fuzzing (deep edge)",
    "E2E TC-DAP-1654-024": "cross-region race condition",
    "E2E TC-DAP-1654-025": "TTL timing (deep edge)",
    "E2E TC-DAP-1654-026": "Catapult 500 failure injection (deep edge)",
    "E2E TC-DAP-1654-028": "idempotency replay (deep edge)",
    "E2E TC-DAP-1654-029": "load/soak test",
}
INVENTED = "_Region-mismatched credential login returns HTTP 409, code=ERROR_REGION_MISMATCH"
UNCHECKABLE = "_System works correctly and logs look fine"
R7_TARGET = "CRA_Tempcreds_Regr_032"
R7_ADDITION = "_Username lookup served via gsiUsrnameIx"


def read_section(ws) -> dict:
    tcs = {}
    for r in range(FIRST_TC_ROW, LAST_TC_ROW + 1):
        values = [ws.cell(r, c).value for c in range(1, 6)]
        tcs[int(str(values[0]).rsplit("_", 1)[1])] = values
    return tcs


def build_before(out: Path) -> list:
    wb = load_workbook(BASE)
    ws = wb[SHEET]
    for merged in [m for m in ws.merged_cells.ranges if m.min_row >= SECTION_FIRST_ROW]:
        ws.unmerge_cells(str(merged))
    ws.delete_rows(SECTION_FIRST_ROW, ws.max_row - SECTION_FIRST_ROW + 1)   # trailing rows only
    cleared = []
    version = wb["Version"]
    for row in version.iter_rows():
        if any("DAP-1654" in str(c.value) for c in row if c.value):
            cleared.append(row[0].row)
            for c in row:
                c.value = None
    wb.save(out)
    return cleared


def build_seeded(out: Path) -> dict:
    wb = load_workbook(BASE)
    ws = wb[SHEET]
    real = read_section(ws)

    def tc(n, scenario=None, steps=None, expected=None):
        _, scen, stay, st, exp = real[n]
        return [scenario or scen, stay, steps or st, expected or exp]

    kept = [125, 126, 127, 128, 130, 131, 132, 133, 134, 135, 136, 137]   # 124 (R2) and 129 (R3) dropped
    plan = []   # (content, source, plant)
    for n in kept:
        content, plant = tc(n), None
        if n == 125:
            content[0] = "Priority: P3 — " + content[0]
            plant = "R6"
        elif n == 133:
            content[3] = content[3] + "\n" + INVENTED
            plant = "R1"
        elif n == 134:
            content[3] = UNCHECKABLE
            plant = "R9"
        plan.append((content, REAL_SOURCES[n], plant))
    duplicate = tc(131)
    duplicate[0] = duplicate[0].replace("(Secrets Manager rejects key)", "(Secrets Manager rejects key) - repeat check")
    plan.append((duplicate, "E2E TC-DAP-1654-010, -011", "R4"))
    plan.append(([
        "DAP-1654 Edge - Empty, whitespace, null and oversized field values are rejected. "
        "Boundary: API Gateway -> temp-creds-auth Lambda",
        "Yes",
        '_Valid x-api-key\n_POST with username "", then "   " (whitespace), then null, then a >10KB username string\n'
        "_POST with a malformed JSON body",
        "_Empty/whitespace/null username or password -> HTTP 400 ERROR_BAD_CREDENTIALS\n"
        "_Oversized/malformed body handled gracefully (400, not 500); no stack trace leaked",
    ], "E2E TC-DAP-1654-023", "R5"))
    assert len(plan) == LAST_TC_ROW - FIRST_TC_ROW + 1

    placed = {}
    for offset, (content, source, plant) in enumerate(plan):
        row, number = FIRST_TC_ROW + offset, FIRST_NUMBER + offset
        tc_id = f"{PREFIX}{number}"
        ws.cell(row, 1).value = tc_id
        for col, value in enumerate(content, start=2):
            ws.cell(row, col).value = value
        placed[tc_id] = (source, plant)

    for cell in ws[HEADER_ROW][:5]:                                   # R8: header style lost
        cell.font = Font(bold=False)
        cell.fill = PatternFill(fill_type=None)

    for row in ws.iter_rows(min_col=1, max_col=5):                    # R7: silent edit of an existing TC
        if row[0].value == R7_TARGET:
            row[4].value = (row[4].value or "") + "\n" + R7_ADDITION
            break
    else:
        sys.exit(f"ERROR: {R7_TARGET} not found")

    wb.save(out)
    return placed


def write_summary(path: Path, title: str, added: dict, updated: dict, exclusions: dict,
                  tickets: str = "none", jira_line: str = "none (E2E-only run)") -> None:
    lines = [
        f"# Change summary — {title}",
        "",
        f"- **Tickets:** {tickets} | **E2E files:** {E2E_FILE}",
        f"- **Jira fetched:** {jira_line}",
        "- **User-stated inputs:** none",
        "",
        "## Changes",
        *[f"- ADDED: {tc_id} — [source: {source}]" for tc_id, source in added.items()],
        *[f"- UPDATED: {tc_id} — {what}" for tc_id, what in updated.items()],
        "",
        "## Exclusions",
        *[f"- {item} — {reason}" for item, reason in exclusions.items()],
        "",
        "## Information gaps",
        "- none",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def build_replay() -> None:
    """Case `replay_jira`: the real DAP-1654 enhancement, now WITH the live Jira ticket.

    Tests the Jira paths the E2E-only cases can't: blind Step 3 (independent fetch),
    reconcile Step 4 (re-fetch) and the timestamp comparison. The enhancement was
    written 2026-06-16 and the ticket was last updated 2026-09-24, so the reviewer
    must raise "ticket changed after enhancement" under Needs your decision.
    """
    folder = HERE / "replay_jira"
    folder.mkdir(exist_ok=True)
    for stale in folder.iterdir():
        if stale.is_file():
            stale.unlink()
    if not (HERE / "control" / SNAPSHOT_NAME).exists():
        sys.exit("ERROR: run the full build first (control snapshot missing)")
    shutil.copy(HERE / "control" / SNAPSHOT_NAME, folder / SNAPSHOT_NAME)
    if not (HERE / "live" / "control.xlsx").exists():
        shutil.copy(BASE, HERE / "live" / "control.xlsx")
    write_summary(
        folder / "change-summary.md", "calibration replay_jira",
        {f"{PREFIX}{n}": s for n, s in REAL_SOURCES.items()}, {}, EXCLUSIONS,
        tickets="DAP-1654",
        jira_line="DAP-1654 updated=2026-06-16T11:04:00.000-0400 "
                  "(stand-in: the 2026-06-16 enhancement predates timestamp recording; "
                  "the ticket was edited 2026-09-24)",
    )
    print("built: replay_jira")


def main() -> None:
    if len(sys.argv) > 1 and sys.argv[1] == "--replay-only":
        build_replay()
        return
    if not BASE.exists():
        shutil.copy2(LIVE_SUITE, BASE)
        print(f"froze base suite: {BASE}")
    for folder in ("seeded", "control", "live"):
        (HERE / folder).mkdir(exist_ok=True)
        for stale in (HERE / folder).iterdir():
            if stale.is_file():
                stale.unlink()

    cleared = build_before(HERE / "seeded" / SNAPSHOT_NAME)
    shutil.copy2(HERE / "seeded" / SNAPSHOT_NAME, HERE / "control" / SNAPSHOT_NAME)
    # Plain copy (fresh timestamp): copy2 kept base's 2026-06-16 mtime, which made the
    # live file look older than its snapshot and drew a reviewer query (2026-10-02 run, D-05).
    shutil.copy(BASE, HERE / "live" / "control.xlsx")
    placed = build_seeded(HERE / "live" / "seeded.xlsx")

    write_summary(
        HERE / "control" / "change-summary.md", "calibration control",
        {f"{PREFIX}{n}": s for n, s in REAL_SOURCES.items()}, {}, EXCLUSIONS,
    )
    seeded_added = {tc_id: source for tc_id, (source, _) in placed.items()}
    seeded_added[f"{PREFIX}138"] = "E2E TC-DAP-1654-008"                   # R10: phantom claim
    write_summary(
        HERE / "seeded" / "change-summary.md", "calibration seeded",
        seeded_added, {R7_TARGET: "clarified lookup [reason: E2E TC-DAP-1654-001]"},
        {k: v for k, v in EXCLUSIONS.items() if k != "E2E TC-DAP-1654-023"},  # fuzz TC included instead
    )

    build_replay()
    plants = {plant: tc_id for tc_id, (_, plant) in placed.items() if plant}
    write_answer_key(plants, cleared)
    print("built: seeded, control, live, answer-key.md")
    for check in sorted(plants, key=lambda c: int(c[1:])):
        print(f"  {check}: {plants[check]}")


def write_answer_key(plants: dict, cleared_version_rows: list) -> None:
    key = f"""# Calibration answer key — enhancer review

Generated by `build_calibration.py` from `base_suite.xlsx` (frozen copy of the live suite).
Source for both cases: `{E2E_FILE}` only (tickets: none).
Version-sheet rows cleared in "before" (mentioned DAP-1654): {cleared_version_rows or 'none'}.

## Case `seeded` — one planted mistake per check

| Check | Planted mistake | Where | Expected severity |
|---|---|---|---|
| R1 | Invented error code `ERROR_REGION_MISMATCH` / HTTP 409 (in no source) | {plants['R1']} Expected Result | Critical |
| R2 | Core happy path (E2E TC-DAP-1654-001) not added and not excluded | — (missing) | Major |
| R3 | Regional-failover boundary Route 53 → API Gateway (E2E TC-DAP-1654-008) not added and not excluded | — (missing) | Major |
| R4 | Near-duplicate of the bad-API-key TC (the other copy is {plants['R4']}'s twin) | {plants['R4']} | Major |
| R5 | Input-fuzzing E2E TC-DAP-1654-023 brought into regression (regression rules exclude it) | {plants['R5']} | Minor (Major acceptable) |
| R6 | Integration-boundary TC (E2E P1) marked `Priority: P3` | {plants['R6']} Scenario | Minor (Major acceptable) |
| R7 | Existing `{R7_TARGET}` edited without an `[UPDATED by …]` notice | {R7_TARGET} Expected Result | Minor |
| R8 | New section header row (row {HEADER_ROW}) not bold, no `D9E1F2` fill | row {HEADER_ROW} | Critical |
| R9 | Expected result can't be checked: "{UNCHECKABLE[1:]}" | {plants['R9']} | Minor (Major acceptable) |
| R10 | Summary claims `ADDED: {PREFIX}138`, which doesn't exist | change-summary.md | Major |

**Expected verdict:** FAIL — roll back recommended (R1 + R8 Critical).

**Pass rule (seeded):**
1. All 10 plants are reported, under the right check, with evidence.
2. Severities are as above (or the stated acceptable alternative).
3. The verdict is as above.
4. **No false alarms.** Every other finding must either:
   - also appear in the `control` review (a property of the real DAP-1654 section), **or**
   - be a genuine side effect of the planting. For example, renumbering TCs breaks a cross-reference such as "links to TC 129" — that's a true positive.

   A finding that's none of these is a false alarm.

## Case `control` — the real DAP-1654 enhancement (no plants)

`live\\control.xlsx` is the suite as it actually is. The 14 DAP-1654 TCs (124–137) are what the enhancer really wrote on 2026-06-16, with its recorded exclusions.

This is **not** guaranteed clean: it's real work. Its findings form the **baseline**.
- **Every control finding must be judged by hand:** genuine issue in the real suite (keep it, and fix the suite separately) or false alarm (counts against the reviewer).
- **Pass rule (control):** zero false alarms, and the verdict follows from the genuine findings.

## How to run

1. **Rebuild:** `python build_calibration.py`.
2. **Blind:** run `enhancer-review-blind` **once**, on `seeded`, with
   ```
   run_dir: {HERE}\\seeded
   tickets: none
   e2e_files: {E2E_FILE}
   cycle: 1
   ```
   Then copy `seeded\\expected-coverage.md` to `control\\`. Both cases have the same before-state and sources, so they share one blind pass; this also removes blind-pass variation from the comparison.
3. **Reconcile:** run `enhancer-review-reconcile` on each case:
   - `seeded` with `suite: {HERE}\\live\\seeded.xlsx`;
   - `control` with `suite: {HERE}\\live\\control.xlsx`.
4. **Score** with this key.
"""
    (HERE / "answer-key.md").write_text(key, encoding="utf-8")


if __name__ == "__main__":
    main()
